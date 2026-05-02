from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ACCENT = "2563EB"
ACCENT_DARK = "1E3A8A"
HEADER_BG = "EAF2FF"
SUBTLE_BG = "F8FAFC"
BORDER = "CBD5E1"
TEXT = RGBColor(15, 23, 42)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=10.5, bold=False, color=TEXT) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text_with_bold(paragraph, text: str, size=10.5, color=TEXT) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        set_font(run, size=size, bold=bold, color=color)


def clean_inline(text: str) -> str:
    return text.replace("`", "").strip()


def is_separator(line: str) -> bool:
    value = line.strip().strip("|")
    return bool(value) and all(part.strip().replace(":", "").replace("-", "") == "" for part in value.split("|"))


def is_table_row(line: str) -> bool:
    value = line.strip()
    return "|" in value and not is_separator(value)


def is_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    separator = lines[index + 1].strip()
    return is_table_row(current) and is_separator(separator)


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = index
    while i < len(lines):
        line = lines[i].strip()
        if is_separator(line):
            i += 1
            continue
        if not is_table_row(line):
            break
        rows.append([clean_inline(cell) for cell in line.strip("|").split("|")])
        i += 1
    return rows, i


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_table_borders(table)

    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            value = row[c_idx] if c_idx < len(row) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, HEADER_BG)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, SUBTLE_BG)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.18
            add_text_with_bold(para, value, size=9.2 if cols >= 5 else 9.8, color=TEXT)
            if r_idx == 0:
                for run in para.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 58, 138)

    set_repeat_table_header(table.rows[0])
    document.add_paragraph()


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading("", level=min(level, 3))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(12 if level == 2 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    if level == 1:
        set_font(run, size=22, bold=True, color=RGBColor(30, 58, 138))
    elif level == 2:
        set_font(run, size=15, bold=True, color=RGBColor(37, 99, 235))
    else:
        set_font(run, size=12.5, bold=True, color=RGBColor(51, 65, 85))


def add_bullet(document: Document, text: str, ordered=False) -> None:
    paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.25
    add_text_with_bold(paragraph, clean_inline(text), size=10.2)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.3
    add_text_with_bold(paragraph, clean_inline(text), size=10.5)


def build_docx(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = TEXT

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("晚辅托管系统 MVP 版本需求分析说明书")
    set_font(run, size=8.5, color=RGBColor(100, 116, 139))

    i = 0
    first_title = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "---":
            document.add_paragraph()
            i += 1
            continue
        if is_table(lines, i):
            rows, i = parse_table(lines, i)
            add_table(document, rows)
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            title = clean_inline(heading_match.group(2))
            add_heading(document, title, level)
            if first_title and level == 1:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run("MVP V1.0 | 2026-05-01 | 需求基准文档")
                set_font(r, size=10.5, color=RGBColor(71, 85, 105))
                document.add_section(WD_SECTION.NEW_PAGE)
                first_title = False
            i += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.+)$", line)
        bullet_match = re.match(r"^[-*]\s+(.+)$", line)
        if numbered_match:
            add_bullet(document, numbered_match.group(1), ordered=True)
        elif bullet_match:
            add_bullet(document, bullet_match.group(1), ordered=False)
        else:
            add_paragraph(document, line)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: markdown_to_docx.py input.md output.docx", file=sys.stderr)
        return 2
    build_docx(Path(sys.argv[1]), Path(sys.argv[2]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
